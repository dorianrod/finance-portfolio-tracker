#!/usr/bin/env python3
"""Block commits/pushes that mention personal portfolio data.

This repo is public but `data/` (a symlink to the real portfolio data)
is gitignored and must never leak into code, comments, commit messages
or test fixtures. Rather than maintaining a manual blocklist, this
script derives one straight from `data/`: every distinctive asset name,
ISIN and ticker it finds there is a value that must never appear in
anything pushed to origin.

Reads the text to scan from stdin (a commit message, or `git log -p`
output for a push range) and exits non-zero if any blocklisted token
is found, printing the offending tokens.
"""

import csv
import re
import sys
from pathlib import Path


def repo_root() -> Path:
    return Path(__file__).resolve().parent.parent


def _first_column(index: dict[str, int], *names: str) -> int | None:
    return next((index[name] for name in names if name in index), None)


def _add_csv_tokens(
    identifiers: set[str], names: set[str], path: Path
) -> None:
    try:
        csv_file = path.open(newline="", errors="ignore")
    except OSError:
        return

    with csv_file:
        for row in csv.DictReader(csv_file):
            for header, raw_value in row.items():
                normalized_header = (header or "").strip().lower()
                if raw_value is None:
                    continue
                value = raw_value.strip()
                if not value:
                    continue
                if normalized_header in {
                    "id",
                    "isin",
                    "ticker",
                    "yahoo",
                    "yahoo_symbol",
                }:
                    identifiers.add(value)
                elif normalized_header in {"name", "nom", "nom_placement"}:
                    names.add(value)


def load_blocklist(data_dir: Path) -> set[str]:
    identifiers: set[str] = set()
    names: set[str] = set()

    # data/input/allocations/*.xlsx: asset name, ISIN and ticker columns.
    alloc_dir = data_dir / "input" / "allocations"
    if alloc_dir.exists():
        try:
            import openpyxl
        except ImportError:
            openpyxl = None
        if openpyxl is not None:
            for f in sorted(alloc_dir.glob("*.xlsx")):
                try:
                    wb = openpyxl.load_workbook(
                        f, read_only=True, data_only=True
                    )
                    ws = wb[wb.sheetnames[0]]
                    rows = ws.iter_rows(values_only=True)
                    next(rows)  # title row
                    header = [str(h or "").strip().lower() for h in next(rows)]
                    idx = {h: i for i, h in enumerate(header) if h}
                    id_col = _first_column(idx, "id", "isin")
                    ticker_col = _first_column(idx, "ticker")
                    yahoo_col = _first_column(
                        idx, "yahoo_symbol", "yahoo"
                    )
                    name_col = _first_column(
                        idx, "nom_placement", "name", "nom"
                    )
                    for row in rows:
                        for col in (id_col, ticker_col, yahoo_col):
                            if col is not None and col < len(row) and row[col]:
                                identifiers.add(str(row[col]).strip())
                        if (
                            name_col is not None
                            and name_col < len(row)
                            and row[name_col]
                        ):
                            names.add(str(row[name_col]).strip())
                except Exception:
                    continue

    # data/input/asset_prices/others/*.csv: name,ticker,isin,...
    others_dir = data_dir / "input" / "asset_prices" / "others"
    if others_dir.exists():
        for f in others_dir.glob("*.csv"):
            _add_csv_tokens(identifiers, names, f)

    # generated/*.csv: isin,ticker,yahoo_symbol,name,...
    generated_dir = data_dir / "input" / "asset_prices" / "generated"
    if generated_dir.exists():
        for f in generated_dir.glob("*.csv"):
            _add_csv_tokens(identifiers, names, f)

    # Very short or one-word labels ("Cash", "Bitcoin", ...) are too
    # generic to block safely in source code and commit messages. Instrument
    # names remain distinctive even when their capitalization changes.
    clean_identifiers = {
        token
        for token in identifiers
        if len(token) >= 3 and token.lower() not in ("nan", "none")
    }
    distinctive_names = {
        name
        for name in names
        if len(name) >= 12 and len(name.split()) >= 2
    }
    return clean_identifiers | distinctive_names


def find_matches(text: str, tokens: set[str]) -> set[str]:
    found: set[str] = set()
    for tok in tokens:
        pattern = r"(?<![A-Za-z0-9])" + re.escape(tok) + r"(?![A-Za-z0-9])"
        flags = re.IGNORECASE if any(char.isspace() for char in tok) else 0
        if re.search(pattern, text, flags=flags):
            found.add(tok)
    return found


def main() -> int:
    text = sys.stdin.read()
    tokens = load_blocklist(repo_root() / "data")
    if not tokens:
        # No local data (fresh clone, CI, ...) — nothing to check against.
        return 0

    found = find_matches(text, tokens)
    if found:
        sys.stderr.write(
            "\nBLOCKED: personal portfolio data found in content"
            " being committed/pushed:\n"
        )
        for tok in sorted(found):
            sys.stderr.write(f"  - {tok}\n")
        sys.stderr.write(
            "Remove these from the commit message, code, comments or test"
            " fixtures (use generic placeholders instead), then retry.\n"
            "Bypass (only if this is a false positive): git commit"
            " --no-verify / git push --no-verify\n"
        )
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
